"""Full-data regressions. All large records here are synthetic, not user data."""
import unittest
import tempfile
import time
from pathlib import Path
from unittest.mock import patch
import matplotlib
matplotlib.use('Agg')
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib import rc_context
from chart_dialog import ChartDialog, analyze_series

class Var:
    def __init__(self, value): self.value = value
    def get(self): return self.value
    def set(self, value): self.value = value

def make_dialog(data):
    d = ChartDialog.__new__(ChartDialog)
    d.data = data.copy(); d.lang = 'en'; d.source_name = ''; d.series = {}
    d.figure = Figure(figsize=(6, 4), dpi=70); d.canvas = FigureCanvasAgg(d.figure)
    d.last_stats = {}; d._preview_after = None
    for attr, value in {'title_var':'Synthetic', 'x_axis':'x' if 'x' in data else 'Record index',
                        'x_name':'X', 'chart_type':'Line', 'grid_var':True, 'legend_var':True}.items():
        setattr(d, attr, Var(value))
    for key, color in [('values', '#2563eb'), ('percent', '#f97316')]:
        if key not in data: continue
        parsed, unit = analyze_series(data[key])
        d.series[key] = {'parsed':parsed, 'detected_unit':unit, 'color':color,
                         'enabled':Var(True), 'name':Var(key), 'unit':Var(unit)}
    return d

def synthetic(n):
    x = np.arange(n, dtype=float)
    x[n//2:n//2+4] = [19, 19, 3, 2]
    values = [f'{3500+i%700}mV' for i in range(n)]
    percent = [f'{i%101}%' for i in range(n)]
    values[5] = 'bad'; percent[5] = None
    values[6] = '4.17V'; values[7] = '2A'; values[8] = '1e999mV'
    values[0] = '-9999mV'; values[-1] = '9999mV'
    values[12] = '1.2345mV'; values[n//2] = '12.75mV'; percent[n//2] = '25.5%'
    x[10] = np.inf; x[11] = np.nan
    return pd.DataFrame({'x': x, 'values': values, 'percent': percent}, index=np.arange(n)+100)

class FullChartTests(unittest.TestCase):
    def test_full_line_scatter_markers(self):
        for n in (73181, 80645):
            d = make_dialog(synthetic(n))
            chosen, rows = d.plotting_data()
            self.assertEqual(len(rows), n)
            self.assertEqual([r[2] for r in rows], list(range(n)))
            self.assertEqual(rows[5][1], [None, None])
            self.assertEqual(rows[6][1][0], 4170.0)
            self.assertIsNone(rows[7][1][0])
            for chart_type in ('Line', 'Scatter', 'Line + markers'):
                with self.subTest(n=n, chart_type=chart_type):
                    d.chart_type.set(chart_type); started = time.perf_counter()
                    self.assertEqual(d.draw(), (n-3, 3))
                    self.assertEqual(d.last_stats['included'], n)
                    self.assertEqual(d.last_stats['series'], {'values':n-5, 'percent':n-3})
                    for i, ax in enumerate(d.figure.axes):
                        y = np.array([np.nan if r[1][i] is None else r[1][i] for r in rows])
                        x = d.data.x.to_numpy().copy(); x[~np.isfinite(x)] = np.nan
                        if chart_type == 'Scatter':
                            offsets = ax.collections[0].get_offsets()
                            self.assertEqual(offsets.shape, (n, 2))
                            mask = ~np.isfinite(x) | ~np.isfinite(y)
                            np.testing.assert_array_equal(np.ma.getmaskarray(offsets)[:,0], mask)
                            np.testing.assert_allclose(offsets[:,0].compressed(), x[~mask])
                            np.testing.assert_allclose(offsets[:,1].compressed(), y[~mask])
                        else:
                            line = ax.lines[0]
                            self.assertEqual(len(line.get_xdata()), n)
                            np.testing.assert_allclose(line.get_xdata(), x, equal_nan=True)
                            np.testing.assert_allclose(line.get_ydata(), y, equal_nan=True)
                            self.assertIsNone(line.get_markevery())
                            self.assertFalse(line.get_path().should_simplify)
                            self.assertFalse(line._subslice)
                    d.canvas.draw()
                    if chart_type != 'Scatter':
                        with rc_context({'path.simplify':True, 'agg.path.chunksize':100}):
                            for ax in d.figure.axes:
                                line = ax.lines[0]
                                line.set_data(line.get_xdata(), line.get_ydata())
                                line.recache(always=True)
                                self.assertFalse(line.get_path().should_simplify)
                            d.canvas.draw()
                            for ax in d.figure.axes:
                                self.assertFalse(ax.lines[0].get_path().should_simplify)
                    print(f'FULL_DRAW n={n} type={chart_type} seconds={time.perf_counter()-started:.3f}', flush=True)

    def test_full_excel_each_cell(self):
        for n in (73181, 80645):
            d = make_dialog(synthetic(n)); started = time.perf_counter()
            _, expected = d.plotting_data()
            with tempfile.TemporaryDirectory() as temp:
                out = Path(temp) / 'synthetic.xlsx'; d.export_excel(out)
                wb = load_workbook(out, read_only=True, data_only=True)
                ws = wb['Chart Data']; self.assertEqual(ws.max_row, n+1)
                actual = ws.iter_rows(values_only=True)
                self.assertEqual(next(actual), ('X', 'values（mV）', 'percent（%）', 'Record idx (0-based)'))
                for position, (x, values, idx) in enumerate(expected):
                    if pd.isna(x) or not np.isfinite(x): x = None
                    self.assertEqual(next(actual), (x, *values, idx), position)
                self.assertIsNone(next(actual, None)); wb.close()
            print(f'FULL_XLSX n={n} seconds={time.perf_counter()-started:.3f}', flush=True)

    def test_missing_x_and_categories(self):
        cases = [(['1','bad','inf',None,'2'], [1,np.nan,np.nan,np.nan,2], 2),
                 (['a','b','a',None,'c'], [0,1,0,np.nan,2], 4),
                 (['2026-10-28','bad',None,'2026-10-28'], None, 2)]
        for xs, expected, valid in cases:
            d = make_dialog(pd.DataFrame({'x':xs, 'values':['1V']*len(xs)}))
            self.assertEqual(d.draw(), (valid,len(xs)-valid))
            line = d.figure.axes[0].lines[0]
            self.assertEqual(len(line.get_xdata()), len(xs))
            self.assertEqual(sum(d.last_stats['x_invalid'].values()), len(xs)-valid)
            if expected is not None:
                np.testing.assert_allclose(line.get_xdata(), expected, equal_nan=True)

    def test_no_series_and_all_invalid_clear_old_chart(self):
        d = make_dialog(pd.DataFrame({'values':['1V','2V']})); d.draw()
        d.series['values']['enabled'].set(False)
        self.assertEqual(len(d.plotting_data()[1]), 2)
        with self.assertRaisesRegex(ValueError, 'at least one'): d.draw()
        self.assertEqual(d.figure.axes, [])
        with patch('chart_dialog.messagebox.showwarning') as warning, patch('chart_dialog.filedialog.asksaveasfilename') as dialog:
            d.export_chart(); warning.assert_called_once(); dialog.assert_not_called()
        self.assertEqual(d.figure.axes, [])
        d.series['values']['enabled'].set(True); d.draw()
        d.series['values']['unit'].set('A')
        with self.assertRaisesRegex(ValueError, 'No valid coordinates'): d.draw()
        self.assertEqual(d.figure.axes, [])
        self.assertEqual(d.last_stats['valid'], 0)
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)/'existing.xlsx'; out.write_bytes(b'unchanged')
            with self.assertRaises(ValueError): d.export_excel(out)
            self.assertEqual(out.read_bytes(), b'unchanged')
            d.data = range(1048576)
            with self.assertRaisesRegex(ValueError, 'row limit'): d.export_excel(out)
            self.assertEqual(out.read_bytes(), b'unchanged')

    def test_custom_unit_label_is_not_conversion(self):
        from chart_dialog import convert_series
        self.assertEqual(convert_series([(4.0,'V'),(4000.0,'mV')], 'V', 'Sensor units'), [4.0,4000.0])
        self.assertEqual(convert_series([(4.0,'V'),(2.0,'A')], 'V', 'Sensor units'), [4.0,None])
        overflow = make_dialog(pd.DataFrame({'values':['1e308kV','1kV']}))
        overflow.series['values']['unit'].set('uV')
        self.assertEqual(overflow.draw(), (1,1))
        self.assertEqual(overflow.last_stats['incompatible_units']['values'], 0)
        d = make_dialog(pd.DataFrame({'values':['4V','5V']}))
        d.series['values']['unit'].set('Sensor units'); d.draw()
        self.assertIn('label only; no conversion', d._stats_text())

    def test_long_integer_and_numeric_text_x_are_lossless(self):
        base = 9007199254740992
        for xs in ([base,base+1,base+2,base+1],
                   [str(base),str(base+1),str(base+2),str(base+1)],
                   ['001','002','001','003']):
            d = make_dialog(pd.DataFrame({'x':xs, 'values':['1V']*4}))
            self.assertEqual(d.draw(), (4,0))
            self.assertTrue(d._text_x)
            coords = d.figure.axes[0].lines[0].get_xdata()
            unique = list(dict.fromkeys(str(x) for x in xs))
            np.testing.assert_array_equal(coords, [unique.index(str(x)) for x in xs])
            self.assertEqual([tick.get_text() for tick in d.figure.axes[0].get_xticklabels()], unique)
            self.assertEqual([r[0] for r in d.plotting_data()[1]], xs)
            with tempfile.TemporaryDirectory() as temp:
                target = Path(temp)/'ids.xlsx'; d.export_excel(target)
                wb = load_workbook(target); rows = list(wb['Chart Data'].iter_rows(min_row=2))
                self.assertEqual([[cell.value for cell in row] for row in rows],
                                 [[str(x),1.0,i] for i,x in enumerate(xs)])
                self.assertTrue(all(row[0].data_type == 's' for row in rows)); wb.close()
        ordinary = make_dialog(pd.DataFrame({'x':[1,2,1,3], 'values':['1V']*4}))
        ordinary.draw(); self.assertFalse(ordinary._text_x)


    def test_empty_unit_hides_label_not_conversion_or_dimension(self):
        from chart_dialog import convert_series
        self.assertEqual(convert_series([(1.0,'V'),(2.0,'A'),(3.0,'V'),(1000.0,'mV')], 'V', ''), [1.0,None,3.0,1.0])
        d = make_dialog(pd.DataFrame({'values':['1V','2A','3V','1000mV']}))
        d.series['values']['unit'].set('')
        self.assertEqual(d.draw(), (3,1))
        self.assertEqual(d.last_stats['incompatible_units']['values'], 1)
        np.testing.assert_allclose(d.figure.axes[0].lines[0].get_ydata(), [1,np.nan,3,1], equal_nan=True)
        self.assertEqual(d.axis_label(d.series['values']), 'values')

    def test_image_exports_keep_all_vertices(self):
        d = make_dialog(synthetic(73181)); d.chart_type.set('Line + markers')
        with tempfile.TemporaryDirectory() as temp:
            for extension in ('png','pdf','svg'):
                target = Path(temp)/('synthetic.'+extension)
                with patch('chart_dialog.filedialog.asksaveasfilename', return_value=str(target)), patch('chart_dialog.messagebox.showinfo'), patch('chart_dialog.messagebox.showerror') as error:
                    d.export_chart(); error.assert_not_called()
                self.assertGreater(target.stat().st_size, 1000)
                for ax in d.figure.axes:
                    self.assertEqual(len(ax.lines[0].get_path().vertices), 73181)
                    self.assertFalse(ax.lines[0].get_path().should_simplify)
                print(f'FULL_IMAGE {extension} bytes={target.stat().st_size}', flush=True)

    def test_nan_gaps_are_not_bridged(self):
        d = make_dialog(pd.DataFrame({'values':['1V',None,'3V','2A','5V']})); d.draw()
        line = d.figure.axes[0].lines[0]
        self.assertTrue(np.isnan(line.get_ydata()[1]))
        self.assertTrue(np.isnan(line.get_ydata()[3]))
        segments = list(line.get_path().iter_segments(remove_nans=True, simplify=False))
        self.assertEqual([code for _, code in segments], [1,1,1])

if __name__ == '__main__':
    unittest.main(verbosity=2)
