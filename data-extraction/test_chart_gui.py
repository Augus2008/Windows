import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from app import App
from chart_dialog import ChartDialog
app=App();app.update_idletasks();app.update()
data=pd.DataFrame({"values":["4184mV","4.17V","bad"],"percent":["100%","98%","97%"]})
dlg=ChartDialog(app,data,"zh","sample.trc");dlg.update_idletasks();dlg.update()
assert set(dlg.series)=={"values","percent"} and dlg.axis_label(dlg.series["values"])=="数值（mV）"
dlg.x_name.set("采样点");valid,skipped=dlg.draw();assert valid==3 and skipped==0
legend=dlg.figure.axes[0].get_legend();assert legend is not None and legend._loc==9
grid_lines=dlg.figure.axes[0].get_xgridlines();assert grid_lines and grid_lines[0].get_color()=="#94a3b8" and abs(grid_lines[0].get_alpha()-.42)<1e-9
with tempfile.TemporaryDirectory() as d:
    out=Path(d)/"chart.xlsx";dlg.export_excel(out);wb=load_workbook(out)
    ws=wb["Chart Data"];preview=wb["Chart Preview"];assert ws["A1"].value=="采样点" and ws.max_row==4
    assert len(preview._images)==1 and not preview._charts and not ws._charts
    with zipfile.ZipFile(out) as archive:
        names=archive.namelist();assert any(n.startswith("xl/media/") for n in names) and not any(n.startswith("xl/charts/") for n in names)
# Saving modified settings must persist them on the parent app.
with patch("chart_dialog.messagebox.askyesnocancel",return_value=True):dlg.request_close()
assert app.chart_settings["x_name"]=="采样点"
dlg2=ChartDialog(app,data,"zh","sample.trc",app.chart_settings);dlg2.update_idletasks();dlg2.update();assert dlg2.x_name.get()=="采样点"
for _ in range(20): dlg2.schedule_preview()
pending = dlg2._preview_after
assert pending in dlg2.tk.call('after','info')
dlg2.destroy()
assert pending not in app.tk.call('after','info')
print("CHART_GUI_TEST_OK")
# Every synthetic result must be plotted and exported in original order.
large=pd.DataFrame({"values":[f"{3500+i%700}mV" for i in range(80645)],"percent":[f"{i%101}%" for i in range(80645)]})
dlg3=ChartDialog(app,large,"zh","large.trc");dlg3.update_idletasks();dlg3.update();chosen,rows=dlg3.plotting_data();assert len(rows)==80645 and [r[2] for r in rows]==list(range(80645))
with tempfile.TemporaryDirectory() as d:
    out=Path(d)/"large-chart.xlsx";dlg3.export_excel(out);wb=load_workbook(out,read_only=False)
    ws=wb["Chart Data"];preview=wb["Chart Preview"];assert ws.max_row==80646
    for idx, row in enumerate(ws.iter_rows(min_row=2,values_only=True)):
        assert row==(idx+1,float(3500+idx%700),float(idx%101),idx), (idx,row)
    assert len(preview._images)==1 and not preview._charts and not ws._charts
    with zipfile.ZipFile(out) as archive:
        names=archive.namelist();assert any(n.startswith("xl/media/") for n in names) and not any(n.startswith("xl/charts/") for n in names)
dlg3.destroy();print("LARGE_CHART_EXCEL_TEST_OK",len(rows))
english=ChartDialog(app,data,'en','sample.trc');english.update_idletasks();english.draw()
assert english.title()=='Chart Settings & Preview'
assert len(english._application_icon_images)==5
english.refresh_preview()
assert 'Included: 3' in english.stats.cget('text') and 'finite points' in english.stats.cget('text')
for width in (980,1180):
    english.geometry(f'{width}x760');english.update_idletasks();english.update()
    assert english.stats.cget('wraplength')<=english.stats.master.winfo_width()-30
english.destroy();app.destroy();print('CHART_ENGLISH_ICON_RESIZE_OK')
